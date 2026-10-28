from datetime import datetime, timedelta
from io import BytesIO

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from auditoria.integridade import verificar_cadeia
from auditoria.models import RegistroAuditoria
from reservas.models import Reserva
from termos.models import TermoResponsabilidade
from usuarios.models import DocumentoCNH, Usuario
from veiculos.models import Veiculo
from vistorias.models import Vistoria

URL_INDICADORES = '/api/relatorios/indicadores/'
URL_OPCOES = '/api/relatorios/opcoes/'
URL_EXPORTAR = '/api/relatorios/exportar/'
ABAS = ['Resumo', 'Usuários', 'Veículos', 'Reservas', 'Termos', 'CNHs', 'Vistorias', 'Auditoria']


def _quando(ano, mes, dia, hora=10, minuto=0):
    return timezone.make_aware(datetime(ano, mes, dia, hora, minuto))


class BaseRelatorio(APITestCase):
    """
    Cenário:
      - 3 veículos: ABC (disponível), DEF (manutenção), GHI (inativo)
      - 5 reservas: mar/concluída (turno manhã, devolvida com atraso, 1000→1080 km),
        mar/negada, abr/aprovada (turno tarde, pernoite), abr/cancelada, mai/pendente
      - João: CNH válida e termo válido; Maria: CNH pendente e sem termo
    """

    def setUp(self):
        self.admin = Usuario.objects.create_user(
            matricula='rel_admin', nome='Admin Relatório', password='senha123', is_admin=True,
        )
        self.joao = Usuario.objects.create_user(
            matricula='rel_joao', nome='João Ação', password='senha123',
        )
        self.maria = Usuario.objects.create_user(
            matricula='rel_maria', nome='Maria Conceição', password='senha123',
        )
        self.abc = self._veiculo('ABC1D23', 'disponivel')
        self.defv = self._veiculo('DEF4G56', 'manutencao')
        self._veiculo('GHI7J89', 'inativo')

        self.concluida = self._reserva(
            self.joao, self.abc, _quando(2026, 3, 10, 6), 'concluida',
            modalidade='turno', turno='manha', data_fim=_quando(2026, 3, 10, 12),
            data_devolucao_real=_quando(2026, 3, 10, 12, 30),
        )
        self.negada = self._reserva(
            self.joao, self.abc, _quando(2026, 3, 15), 'negada',
            observacao_admin='VEÍCULO RESERVADO PARA OUTRA DEMANDA.',
        )
        self._reserva(
            self.maria, self.defv, _quando(2026, 4, 5, 12), 'aprovada',
            modalidade='turno', turno='tarde', pernoite=True, unidade_pernoite='GAMA',
        )
        self._reserva(self.joao, self.defv, _quando(2026, 4, 10), 'cancelada')
        self._reserva(self.maria, self.abc, _quando(2026, 5, 20), 'pendente')

        self._vistoria(self.concluida, 'checkin', 1000, _quando(2026, 3, 10, 6, 5))
        self._vistoria(
            self.concluida, 'checkout', 1080, _quando(2026, 3, 10, 12, 30),
            observacoes='RISCO NA PORTA TRASEIRA.',
        )

        futuro = timezone.localdate() + timedelta(days=200)
        DocumentoCNH.objects.create(
            usuario=self.joao, arquivo='cnh/joao.pdf', status='aprovado', data_validade=futuro,
        )
        DocumentoCNH.objects.create(usuario=self.maria, arquivo='cnh/maria.pdf')
        TermoResponsabilidade.objects.create(
            usuario=self.joao, entidade='SENAI', cpf='529.982.247-25', cnh='02650306461',
            data_solicitacao=timezone.localdate(), status='aprovado', data_validade=futuro,
        )

    def _veiculo(self, placa, situacao):
        return Veiculo.objects.create(
            placa=placa, modelo='Onix', marca='Chevrolet', ano=2024, cor='Prata',
            tipo_combustivel='flex', capacidade=5, status=situacao,
        )

    def _reserva(self, usuario, veiculo, inicio, situacao, **extra):
        extra.setdefault('data_fim', inicio + timedelta(hours=2))
        return Reserva.objects.create(
            usuario=usuario, veiculo=veiculo, data_inicio=inicio, status=situacao,
            motivo='SERVIÇO EXTERNO', **extra,
        )

    def _vistoria(self, reserva, tipo, km, quando, observacoes=None):
        vistoria = Vistoria.objects.create(
            reserva=reserva, tipo=tipo, quilometragem=km, realizada_por=reserva.usuario,
            unidade='taguatinga', observacoes=observacoes,
        )
        Vistoria.objects.filter(pk=vistoria.pk).update(data_hora=quando)
        return vistoria

    def _indicadores(self, **params):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get(URL_INDICADORES, params)
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        return resp.data

    def _exportar(self, **params):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get(URL_EXPORTAR, params)
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        return resp, load_workbook(BytesIO(resp.content))


class PermissoesRelatorioTests(BaseRelatorio):
    def test_anonimo_nao_acessa(self):
        for url in (URL_INDICADORES, URL_OPCOES, URL_EXPORTAR):
            self.assertEqual(self.client.get(url).status_code, status.HTTP_401_UNAUTHORIZED, url)

    def test_funcionario_recebe_403_e_nada_e_auditado(self):
        self.client.force_authenticate(user=self.joao)
        for url in (URL_INDICADORES, URL_OPCOES, URL_EXPORTAR):
            self.assertEqual(self.client.get(url).status_code, status.HTTP_403_FORBIDDEN, url)
        self.assertFalse(RegistroAuditoria.objects.filter(acao='exportar_dados').exists())

    def test_admin_em_primeiro_acesso_e_bloqueado(self):
        self.admin.precisa_definir_senha = True
        self.admin.save()
        self.client.force_authenticate(user=self.admin)
        self.assertEqual(self.client.get(URL_EXPORTAR).status_code, status.HTTP_403_FORBIDDEN)

    def test_opcoes_de_filtro(self):
        self.client.force_authenticate(user=self.admin)
        dados = self.client.get(URL_OPCOES).data
        self.assertEqual(len(dados['veiculos']), 3)
        self.assertEqual(
            {u['value'] for u in dados['usuarios']}, {self.joao.id, self.maria.id},
        )
        self.assertIn('negada', [s['value'] for s in dados['status']])
        self.assertEqual([t['value'] for t in dados['turnos']], ['manha', 'tarde', 'noite'])


class IndicadoresTests(BaseRelatorio):
    @staticmethod
    def _por(lista, valor):
        return next(item['total'] for item in lista if item['valor'] == valor)

    def test_indicadores_sem_filtros(self):
        dados = self._indicadores()

        reservas = dados['reservas']
        self.assertEqual(reservas['total'], 5)
        self.assertEqual(reservas['pendentes'], 1)
        self.assertEqual(reservas['aprovadas'], 2)
        self.assertEqual(reservas['negadas'], 1)
        self.assertEqual(reservas['canceladas'], 1)
        self.assertEqual(reservas['taxa_aprovacao'], 66.7)
        self.assertEqual(reservas['pernoites'], 1)
        self.assertEqual(reservas['devolucoes_registradas'], 1)
        self.assertEqual(reservas['devolucoes_com_atraso'], 1)
        self.assertEqual(reservas['checkout_atrasado_concluido'], 1)
        self.assertEqual(reservas['checkout_atrasado_aberto'], 0)
        self.assertEqual(reservas['cancelamentos_automaticos'], 0)
        self.assertEqual(reservas['pernoites_validas'], 1)
        self.assertEqual(reservas['falta_checkin_obrigatorio'], 1)
        self.assertEqual(self._por(reservas['por_turno'], 'manha'), 1)
        self.assertEqual(self._por(reservas['por_turno'], 'tarde'), 1)
        self.assertEqual(self._por(reservas['por_turno'], 'noite'), 0)
        self.assertEqual(self._por(reservas['por_turno'], 'horario'), 3)
        self.assertEqual(
            [(m['mes'], m['total']) for m in reservas['por_mes']],
            [('2026-03', 2), ('2026-04', 2), ('2026-05', 1)],
        )
        self.assertEqual(reservas['top_veiculos'][0]['veiculo_id'], self.abc.id)
        self.assertEqual(reservas['top_veiculos'][0]['total'], 3)

        frota = dados['frota']
        self.assertEqual(frota['total'], 3)
        self.assertEqual(frota['ativos'], 2)
        self.assertEqual(frota['disponiveis'], 1)
        self.assertEqual(frota['manutencao'], 1)
        self.assertEqual(frota['inativos'], 1)
        self.assertEqual(frota['percentual_disponivel'], 50.0)
        self.assertEqual(frota['km_rodados_total'], 80)
        self.assertEqual(frota['uso_por_veiculo'][0]['veiculo_id'], self.abc.id)
        self.assertEqual(frota['uso_por_veiculo'][0]['km_rodados'], 80)

        usuarios = dados['usuarios']
        self.assertEqual(usuarios['funcionarios'], 2)
        self.assertEqual(usuarios['administradores'], 1)
        self.assertEqual(self._por(usuarios['cnh_por_situacao'], 'aprovada'), 1)
        self.assertEqual(self._por(usuarios['cnh_por_situacao'], 'pendente'), 1)
        self.assertEqual(self._por(usuarios['termo_por_situacao'], 'valido'), 1)
        self.assertEqual(self._por(usuarios['termo_por_situacao'], 'sem_termo'), 1)
        self.assertEqual(self._por(usuarios['cnh_enviadas_periodo'], 'pendente'), 1)

        vistorias = dados['vistorias']
        self.assertEqual(vistorias['total'], 2)
        self.assertEqual(vistorias['checkins'], 1)
        self.assertEqual(vistorias['checkouts'], 1)
        self.assertEqual(vistorias['com_observacoes'], 1)
        self.assertEqual(vistorias['km_rodados_total'], 80)
        self.assertEqual(vistorias['viagens_com_km'], 1)

    def test_filtro_de_periodo(self):
        dados = self._indicadores(de='2026-04-01', ate='2026-04-30')
        self.assertEqual(dados['reservas']['total'], 2)
        self.assertEqual(dados['vistorias']['total'], 0)
        self.assertEqual(dados['vistorias']['km_rodados_total'], 0)
        self.assertEqual(dados['filtros']['de'], '2026-04-01')

    def test_filtros_de_veiculo_usuario_status_e_turno(self):
        self.assertEqual(self._indicadores(veiculo=self.abc.id)['reservas']['total'], 3)
        self.assertEqual(self._indicadores(veiculo=self.abc.id)['frota']['total'], 1)
        self.assertEqual(self._indicadores(usuario=self.maria.id)['reservas']['total'], 2)
        self.assertEqual(self._indicadores(usuario=self.maria.id)['usuarios']['funcionarios'], 1)
        self.assertEqual(self._indicadores(status='negada')['reservas']['total'], 1)
        manha = self._indicadores(turno='manha')
        self.assertEqual(manha['reservas']['total'], 1)
        self.assertEqual(manha['vistorias']['km_rodados_total'], 80)

    def test_filtros_invalidos_retornam_400(self):
        self.client.force_authenticate(user=self.admin)
        casos = [
            ({'de': '2026-05-01', 'ate': '2026-04-01'}, 'ate'),
            ({'de': '31/12/2026'}, 'de'),
            ({'status': 'inexistente'}, 'status'),
            ({'turno': 'madrugada'}, 'turno'),
            ({'veiculo': 999999}, 'veiculo'),
            ({'usuario': 'abc'}, 'usuario'),
        ]
        for params, campo in casos:
            resp = self.client.get(URL_INDICADORES, params)
            self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST, params)
            self.assertIn(campo, resp.data)

    def test_indicadores_nao_alteram_contas(self):
        antes = list(Usuario.objects.order_by('id').values_list('is_active', flat=True))
        self._indicadores()
        depois = list(Usuario.objects.order_by('id').values_list('is_active', flat=True))
        self.assertEqual(antes, depois)


class ExportacaoExcelTests(BaseRelatorio):
    def _valores(self, ws):
        return [[celula.value for celula in linha] for linha in ws.iter_rows()]

    def test_gera_xlsx_valido_com_todas_as_abas(self):
        resp, wb = self._exportar()
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment;', resp['Content-Disposition'])
        self.assertIn('.xlsx', resp['Content-Disposition'])
        self.assertEqual(wb.sheetnames, ABAS)

        reservas = self._valores(wb['Reservas'])
        self.assertEqual(reservas[0][:4], ['ID', 'Solicitante', 'Matrícula', 'Veículo'])
        self.assertEqual(len(reservas), 1 + 5)
        self.assertIsInstance(reservas[1][6], datetime)
        self.assertEqual(reservas[1][6], datetime(2026, 3, 10, 6, 0))

        nomes = [linha[1] for linha in self._valores(wb['Usuários'])[1:]]
        self.assertIn('João Ação', nomes)
        self.assertIn('Maria Conceição', nomes)
        self.assertEqual(len(self._valores(wb['Veículos'])), 1 + 3)
        self.assertEqual(len(self._valores(wb['Vistorias'])), 1 + 2)
        self.assertEqual(len(self._valores(wb['CNHs'])), 1 + 2)
        self.assertEqual(len(self._valores(wb['Termos'])), 1 + 1)

    def test_nao_exporta_segredos_nem_documentos_pessoais(self):
        _, wb = self._exportar()
        textos = ' '.join(
            str(valor) for ws in wb.worksheets for linha in self._valores(ws) for valor in linha
            if valor is not None
        )
        for usuario in (self.admin, self.joao, self.maria):
            self.assertNotIn(usuario.password, textos)
        self.assertNotIn('pbkdf2', textos)
        self.assertNotIn('529.982.247-25', textos)
        self.assertNotIn('02650306461', textos)
        self.assertNotIn('cnh/joao.pdf', textos)

    def test_exportacao_respeita_filtros(self):
        _, wb = self._exportar(status='negada', usuario=self.joao.id)
        reservas = self._valores(wb['Reservas'])
        self.assertEqual(len(reservas), 1 + 1)
        self.assertEqual(reservas[1][0], self.negada.id)
        self.assertEqual(reservas[1][17], 'VEÍCULO RESERVADO PARA OUTRA DEMANDA.')
        self.assertEqual(len(self._valores(wb['Usuários'])), 1 + 1)

        _, wb = self._exportar(de='2026-04-01', ate='2026-04-30')
        self.assertEqual(len(self._valores(wb['Reservas'])), 1 + 2)
        self.assertEqual(len(self._valores(wb['Vistorias'])), 1)

        resumo = {linha[0]: linha[1] for linha in self._valores(wb['Resumo']) if linha[0]}
        self.assertEqual(resumo['Data inicial'], '2026-04-01')
        self.assertEqual(resumo['Veículo'], 'Todos')

    def test_exportacao_e_auditada_com_filtros(self):
        self._exportar(de='2026-03-01', veiculo=self.abc.id)
        log = RegistroAuditoria.objects.get(acao='exportar_dados')
        self.assertEqual(log.entidade, 'relatorio')
        self.assertEqual(log.ator_id, self.admin.id)
        self.assertEqual(log.ator_nome, 'Admin Relatório')
        self.assertEqual(log.detalhes['formato'], 'xlsx')
        self.assertEqual(log.detalhes['filtros']['de'], '2026-03-01')
        self.assertEqual(log.detalhes['filtros']['veiculo'], 'ABC1D23')
        self.assertEqual(log.detalhes['linhas']['Reservas'], 3)
        self.assertTrue(verificar_cadeia()['integra'])

    def test_filtro_invalido_nao_gera_arquivo_nem_auditoria(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get(URL_EXPORTAR, {'de': '2026-05-01', 'ate': '2026-01-01'})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(RegistroAuditoria.objects.filter(acao='exportar_dados').exists())

    def test_texto_com_formula_e_gravado_como_texto(self):
        self._reserva(
            self.joao, self.abc, _quando(2026, 6, 1), 'pendente',
            destino='=HYPERLINK("http://exemplo.invalido")',
        )
        _, wb = self._exportar(de='2026-06-01')
        celula = wb['Reservas'].cell(row=2, column=11)
        self.assertEqual(celula.value, '=HYPERLINK("http://exemplo.invalido")')
        self.assertEqual(celula.data_type, 's')

    def test_datas_de_validade_sao_datas(self):
        _, wb = self._exportar()
        termos = self._valores(wb['Termos'])
        self.assertIsInstance(termos[1][9], datetime)
        self.assertEqual(termos[1][9].date(), timezone.localdate() + timedelta(days=200))
        self.assertEqual(wb['Termos'].cell(row=2, column=10).number_format, 'DD/MM/YYYY')
