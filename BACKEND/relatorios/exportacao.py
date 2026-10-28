"""
Planilha administrativa (.xlsx) gerada com openpyxl.

Não entram na planilha: senhas/hashes de senha, tokens, arquivos ou links de
arquivos, CPF e número da CNH declarados no termo, RENAVAM/chassi e
geolocalização das vistorias.
"""
from datetime import date, datetime
from io import BytesIO

from django.db.models import Count
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from auditoria.models import RegistroAuditoria
from reservas.models import Reserva
from termos.models import TermoResponsabilidade
from usuarios.models import DocumentoCNH, Usuario
from veiculos.models import Veiculo
from vistorias.models import Vistoria

from .indicadores import (
    STATUS_APROVADOS,
    SITUACOES_CNH,
    SITUACOES_TERMO,
    calcular_indicadores,
    km_por_reserva,
    situacao_documentos_usuarios,
)

FORMATO_DATA = 'DD/MM/YYYY'
FORMATO_DATA_HORA = 'DD/MM/YYYY HH:MM'
LARGURA_MAXIMA = 60
COR_CABECALHO = PatternFill('solid', fgColor='1D3047')
FONTE_CABECALHO = Font(bold=True, color='FFFFFF')

ROTULOS_CNH = dict(SITUACOES_CNH)
ROTULOS_TERMO = dict(SITUACOES_TERMO)
ROTULOS_FILTROS = {
    'de': 'Data inicial',
    'ate': 'Data final',
    'veiculo': 'Veículo',
    'usuario': 'Usuário',
    'status': 'Status da reserva',
    'turno': 'Turno',
}


def _valor_celula(valor):
    """Datas com fuso viram horário local sem tzinfo (o Excel não guarda fuso)."""
    if isinstance(valor, datetime):
        if timezone.is_aware(valor):
            valor = timezone.localtime(valor)
        return valor.replace(tzinfo=None, microsecond=0)
    if isinstance(valor, bool):
        return 'Sim' if valor else 'Não'
    return valor


def _escrever(ws, linha, coluna, valor):
    celula = ws.cell(row=linha, column=coluna, value=_valor_celula(valor))
    # Texto iniciado por "=" seria gravado como fórmula; força texto literal.
    if celula.data_type == 'f':
        celula.data_type = 's'
    if isinstance(valor, datetime):
        celula.number_format = FORMATO_DATA_HORA
    elif isinstance(valor, date):
        celula.number_format = FORMATO_DATA
    return celula


def _aba(wb, titulo, cabecalhos, linhas):
    ws = wb.create_sheet(titulo)
    larguras = [len(c) for c in cabecalhos]
    for coluna, texto in enumerate(cabecalhos, start=1):
        celula = ws.cell(row=1, column=coluna, value=texto)
        celula.font = FONTE_CABECALHO
        celula.fill = COR_CABECALHO
        celula.alignment = Alignment(vertical='center')
    total = 0
    for total, valores in enumerate(linhas, start=1):
        for coluna, valor in enumerate(valores, start=1):
            _escrever(ws, total + 1, coluna, valor)
            tamanho = 16 if isinstance(valor, (date, datetime)) else len(str(valor or ''))
            larguras[coluna - 1] = max(larguras[coluna - 1], tamanho)
    for coluna, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(coluna)].width = min(largura + 2, LARGURA_MAXIMA)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cabecalhos))}{max(total + 1, 1)}'
    return total


def _aba_resumo(ws, indicadores, usuario, contagens):
    ws.title = 'Resumo'
    ws['A1'] = 'Relatório administrativo — Frota SENAI'
    ws['A1'].font = Font(bold=True, size=14)
    linhas = [
        ('Gerado em', timezone.now()),
        ('Gerado por', f'{usuario.nome} ({usuario.matricula})'),
        (None, None),
        ('Filtros aplicados', None),
    ]
    for chave, rotulo in ROTULOS_FILTROS.items():
        linhas.append((rotulo, indicadores['filtros'].get(chave) or 'Todos'))

    frota, reservas = indicadores['frota'], indicadores['reservas']
    usuarios, vistorias = indicadores['usuarios'], indicadores['vistorias']
    taxa = reservas['taxa_aprovacao']
    linhas += [
        (None, None),
        ('Frota (situação atual)', None),
        ('Veículos cadastrados', frota['total']),
        ('Veículos ativos', frota['ativos']),
        ('Disponíveis', frota['disponiveis']),
        ('Em uso', frota['em_uso']),
        ('Em manutenção', frota['manutencao']),
        ('Inativos', frota['inativos']),
        (None, None),
        ('Reservas no período', None),
        ('Total', reservas['total']),
        ('Pendentes', reservas['pendentes']),
        ('Aprovadas (inclui em andamento e concluídas)', reservas['aprovadas']),
        ('Negadas', reservas['negadas']),
        ('Canceladas', reservas['canceladas']),
        ('Taxa de aprovação (%)', taxa if taxa is not None else '—'),
        ('Com pernoite (válidas, não dispensam check-in)', reservas['pernoites_validas']),
        ('Cancelamentos automáticos por falta de check-in', reservas['cancelamentos_automaticos']),
        ('Falta de check-in obrigatório', reservas['falta_checkin_obrigatorio']),
        ('Check-outs atrasados concluídos', reservas['checkout_atrasado_concluido']),
        ('Check-outs atrasados em aberto', reservas['checkout_atrasado_aberto']),
        ('Devoluções com atraso', reservas['devolucoes_com_atraso']),
        (None, None),
        ('Usuários (situação atual)', None),
        ('Funcionários', usuarios['funcionarios']),
        ('Funcionários ativos', usuarios['funcionarios_ativos']),
        ('Administradores', usuarios['administradores']),
    ]
    linhas += [(f'CNH — {i["rotulo"]}', i['total']) for i in usuarios['cnh_por_situacao']]
    linhas += [(f'Termo — {i["rotulo"]}', i['total']) for i in usuarios['termo_por_situacao']]
    linhas += [
        (None, None),
        ('Vistorias no período', None),
        ('Check-ins', vistorias['checkins']),
        ('Check-outs', vistorias['checkouts']),
        ('Com observações', vistorias['com_observacoes']),
        ('Km rodados (check-out − check-in)', vistorias['km_rodados_total']),
        (None, None),
        ('Linhas exportadas por aba', None),
    ]
    linhas += list(contagens.items())

    for indice, (rotulo, valor) in enumerate(linhas, start=3):
        if rotulo is None:
            continue
        celula = _escrever(ws, indice, 1, rotulo)
        if valor is None:
            celula.font = Font(bold=True)
        else:
            _escrever(ws, indice, 2, valor)
    ws.column_dimensions['A'].width = 48
    ws.column_dimensions['B'].width = 40


def _nome(usuario):
    return usuario.nome if usuario else ''


def gerar_planilha(filtros, usuario):
    """Retorna (bytes do xlsx, contagem de linhas por aba)."""
    funcionarios = filtros.usuarios(Usuario.objects.filter(is_admin=False))
    situacoes = situacao_documentos_usuarios(funcionarios)
    cnh_situacao, termo_situacao = situacoes
    indicadores = calcular_indicadores(filtros, situacoes)

    wb = Workbook()
    resumo = wb.active
    contagens = {}

    usuarios = filtros.usuarios(Usuario.objects.all()).order_by('nome')
    contagens['Usuários'] = _aba(
        wb, 'Usuários',
        ['ID', 'Nome', 'Matrícula', 'E-mail', 'Telefone', 'Setor', 'Perfil', 'Conta ativa',
         'Situação do termo', 'Situação da CNH', 'Validade da CNH', 'Cadastrado em'],
        (
            [
                u.id, u.nome, u.matricula, u.email or '', u.telefone or '', u.setor or '',
                'Administrador' if u.is_admin else 'Funcionário', u.is_active,
                '—' if u.is_admin else ROTULOS_TERMO.get(termo_situacao[u.pk][0], ''),
                '—' if u.is_admin else ROTULOS_CNH.get(cnh_situacao[u.pk]['situacao'], ''),
                None if u.is_admin else cnh_situacao[u.pk]['data_validade'],
                u.criado_em,
            ]
            for u in usuarios
        ),
    )

    reservas = filtros.reservas(Reserva.objects.all())
    efetivas = dict(
        reservas.filter(status__in=STATUS_APROVADOS)
        .order_by().values_list('veiculo_id').annotate(total=Count('id'))
    )
    km_veiculo = {}
    for veiculo_id, km in km_por_reserva(reservas).values():
        km_veiculo[veiculo_id] = km_veiculo.get(veiculo_id, 0) + km
    contagens['Veículos'] = _aba(
        wb, 'Veículos',
        ['ID', 'Placa', 'Marca', 'Modelo', 'Ano', 'Cor', 'Combustível', 'Capacidade', 'Status',
         'Km atual', 'Última manutenção', 'Reservas aprovadas no período', 'Km rodados no período'],
        (
            [
                v.id, v.placa, v.marca, v.modelo, v.ano, v.cor, v.get_tipo_combustivel_display(),
                v.capacidade, v.get_status_display(), v.km_atual, v.data_ultima_manutencao,
                efetivas.get(v.id, 0), km_veiculo.get(v.id, 0),
            ]
            for v in filtros.veiculos(Veiculo.objects.all()).order_by('placa')
        ),
    )

    contagens['Reservas'] = _aba(
        wb, 'Reservas',
        ['ID', 'Solicitante', 'Matrícula', 'Veículo', 'Modalidade', 'Turno', 'Retirada',
         'Devolução prevista', 'Devolução real', 'Status', 'Destino', 'Motivo', 'Passageiros',
         'Pernoite', 'Unidade do pernoite', 'Decidida por', 'Data da decisão',
         'Observação / justificativa', 'Solicitada em',
         'Unidade do check-in', 'Cancelamento automático', 'Motivo do cancelamento',
         'Check-out atrasado (concluído)', 'Check-out atrasado (em aberto)'],
        (
            [
                r.id, r.usuario.nome, r.usuario.matricula, r.veiculo.placa,
                r.get_modalidade_display(), r.get_turno_display() if r.turno else '',
                r.data_inicio, r.data_fim, r.data_devolucao_real, r.get_status_display(),
                r.destino or '', r.motivo, r.passageiros, r.pernoite, r.unidade_pernoite or '',
                _nome(r.admin_aprovador), r.data_aprovacao, r.observacao_admin or '', r.criado_em,
                r.get_unidade_checkin_display() if r.unidade_checkin else '',
                r.cancelamento_automatico,
                r.get_motivo_cancelamento_display() if r.motivo_cancelamento else '',
                bool(
                    r.status == 'concluida'
                    and r.data_devolucao_real
                    and r.data_devolucao_real > r.data_fim
                ),
                bool(r.status == 'em_andamento' and r.data_fim < timezone.now()),
            ]
            for r in reservas.select_related('usuario', 'veiculo', 'admin_aprovador')
            .order_by('data_inicio', 'id')
        ),
    )

    contagens['Termos'] = _aba(
        wb, 'Termos',
        ['ID', 'Usuário', 'Matrícula', 'Entidade', 'Status', 'Data da solicitação', 'Enviado em',
         'Decidido por', 'Data da decisão', 'Válido até', 'Observação / justificativa'],
        (
            [
                t.id, t.usuario.nome, t.usuario.matricula, t.entidade, t.get_status_display(),
                t.data_solicitacao, t.criado_em, _nome(t.admin_aprovador), t.data_decisao,
                t.data_validade, t.observacao_admin or '',
            ]
            for t in filtros.documentos(TermoResponsabilidade.objects.all())
            .select_related('usuario', 'admin_aprovador').order_by('criado_em', 'id')
        ),
    )

    contagens['CNHs'] = _aba(
        wb, 'CNHs',
        ['ID', 'Usuário', 'Matrícula', 'Status', 'Validade', 'Enviada em', 'Decidida por',
         'Data da decisão', 'Observação / justificativa'],
        (
            [
                d.id, d.usuario.nome, d.usuario.matricula, d.get_status_display(),
                d.data_validade, d.criado_em, _nome(d.admin_aprovador), d.data_decisao,
                d.observacao_admin or '',
            ]
            for d in filtros.documentos(DocumentoCNH.objects.all())
            .select_related('usuario', 'admin_aprovador').order_by('criado_em', 'id')
        ),
    )

    contagens['Vistorias'] = _aba(
        wb, 'Vistorias',
        ['ID', 'Reserva', 'Veículo', 'Solicitante', 'Tipo', 'Unidade', 'Quilometragem',
         'Realizada por', 'Data/hora', 'Fotos', 'Observações'],
        (
            [
                v.id, v.reserva_id, v.reserva.veiculo.placa, v.reserva.usuario.nome,
                v.get_tipo_display(), v.get_unidade_display() if v.unidade else '',
                v.quilometragem, v.realizada_por.nome, v.data_hora, v.total_fotos,
                v.observacoes or '',
            ]
            for v in filtros.vistorias(Vistoria.objects.all())
            .select_related('reserva__veiculo', 'reserva__usuario', 'realizada_por')
            .annotate(total_fotos=Count('fotos')).order_by('data_hora', 'id')
        ),
    )

    contagens['Auditoria'] = _aba(
        wb, 'Auditoria',
        ['Sequência', 'Data/hora', 'Responsável', 'Matrícula do responsável', 'Ação', 'Entidade',
         'ID do registro', 'Identificação', 'Usuário afetado', 'Matrícula do afetado',
         'Descrição', 'Justificativa', 'Hash'],
        (
            [
                a.sequencia, a.criado_em, a.ator_nome, a.ator_matricula, a.get_acao_display(),
                a.get_entidade_display(), a.entidade_id, a.identificacao,
                a.usuario_afetado_nome, a.usuario_afetado_matricula, a.descricao,
                a.justificativa, a.hash,
            ]
            for a in filtros.auditoria(RegistroAuditoria.objects.all()).order_by('sequencia')
        ),
    )

    _aba_resumo(resumo, indicadores, usuario, contagens)

    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue(), contagens
